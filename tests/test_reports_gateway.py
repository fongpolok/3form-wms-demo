import asyncio
import io
import uuid

import openpyxl
import pytest

from mes import services
from site_gateway.drivers.honeywell_tcp import HoneywellTcpDriver
from site_gateway.outbox.store import OfflineQueue

API = "/api/v1"


@pytest.mark.django_db
def test_excel_has_all_sheets(people, clients):
    _, (batch,) = services.create_work_order(people["pic"], {"product_code": "A", "qty": 10, "due_date": "2026-12-01"})
    w = people["worker"]
    services.start_step(w, "WS-01", batch["batch_id"], uuid.uuid4().hex)
    services.record_scrap(w, "WS-01", 1, "SPILL", uuid.uuid4().hex)
    r = clients["em"].get(f"{API}/reports/production.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Work orders", "Step logs", "Scrap", "Downtime", "Measurements", "Stock movements",
                             "Audit trail"]
    assert wb["Scrap"].max_row == 2
    assert clients["worker"].get(f"{API}/reports/production.xlsx").status_code == 403


def test_outbox_survives_restart(tmp_path):
    q = OfflineQueue(tmp_path / "q.sqlite3")
    q.put("scan", {"event_id": "a"})
    q.put("scan", {"event_id": "b"})
    q2 = OfflineQueue(tmp_path / "q.sqlite3")   # simulate gateway restart
    rows = q2.peek()
    assert [p["event_id"] for _, _, p in rows] == ["a", "b"]
    q2.done(rows[0][0])
    assert q2.size() == 1


def test_honeywell_driver_reads_crlf_frames():
    received, statuses = [], []

    async def scenario():
        async def fake_scanner(reader, writer):
            writer.write(b"B-20261019-0001\r\nWIP-X\r\n\r\npartial")
            await writer.drain()
            writer.close()
        server = await asyncio.start_server(fake_scanner, "127.0.0.1", 0)
        port = server.sockets[0].getsockname()[1]
        cfg = {"id": "SCN-T", "host": "127.0.0.1", "port": port, "mode": "client", "terminator": "\\r\\n",
               "workstation": "WS-01"}
        driver = HoneywellTcpDriver(cfg, lambda d, c: received.append(c), lambda d, s: statuses.append(s))
        reader, writer = await asyncio.open_connection("127.0.0.1", port)
        await driver._read_stream(reader, "test")
        server.close()

    asyncio.run(scenario())
    assert received == ["B-20261019-0001", "WIP-X"]
    assert statuses == ["ok", "down"]
