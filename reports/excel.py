"""
Production data workbook: one sheet per data set, filtered by date range.
Used by the "Export Excel" button and the nightly export command.

Sheets: Work orders, Step logs, Scrap, Downtime, Measurements,
        Stock movements, Audit trail.
"""
import io
import logging
from datetime import datetime, time, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.models import AuditLog
from core import config, mongo
from cyclecount.models import QTYRecorder

logger = logging.getLogger("app")
HEADER_FILL = PatternFill("solid", fgColor="0050CC")   # 3Form blue
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _tz():
    return ZoneInfo(config.get("site.timezone", "Asia/Hong_Kong"))


def _local(dt):
    if dt is None:
        return None
    if isinstance(dt, str):
        return dt
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(_tz()).replace(tzinfo=None)


def day_range(date_from, date_to):
    """Inclusive local dates -> UTC datetimes."""
    start = datetime.combine(date_from, time.min, tzinfo=_tz()).astimezone(timezone.utc)
    end = datetime.combine(date_to, time.max, tzinfo=_tz()).astimezone(timezone.utc)
    return start, end


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    limit = int(config.get("excel_export.max_rows_per_sheet", 200000))
    n = 0
    for row in rows:
        if n >= limit:
            ws.append(["(truncated — narrow the date range)"])
            logger.warning("Sheet %s truncated at %d rows", title, limit)
            break
        ws.append(row)
        n += 1
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 4)
    ws.freeze_panes = "A2"
    return n


def _events(db, tenant, start, end, etype):
    return db.step_events.find({"tenant_openid": tenant, "type": etype, "ts": {"$gte": start, "$lte": end}},
                               {"_id": 0}).sort("ts", 1)


def build_workbook(tenant_openid, date_from, date_to):
    start, end = day_range(date_from, date_to)
    db = mongo.db()
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}

    counts["Work orders"] = _sheet(wb, "Work orders",
        ["WO ID", "Product code", "Product", "Client", "Qty", "Due date", "Ship mode", "Source",
         "Special/Urgent", "Status", "Created by", "Created at"],
        ([w["wo_id"], w["product_code"], w.get("product_name"), w.get("client"), w["qty"], w["due_date"],
          w["ship_mode"], w["source"], w["special_urgent"], w["status"], w.get("created_by"), _local(w["created_at"])]
         for w in db.work_orders.find({"tenant_openid": tenant_openid, "created_at": {"$gte": start, "$lte": end}},
                                      {"_id": 0}).sort("created_at", 1)))

    counts["Step logs"] = _sheet(wb, "Step logs",
        ["WIP ID", "Batch ID", "WO ID", "Step", "Step code", "Workstation", "Qty in", "Good qty", "Scrap qty",
         "Downtime (s)", "Started", "Finished", "Started by", "Edited by PIC"],
        ([w["wip_id"], w["batch_id"], w["wo_id"], w["step_no"], w["step_code"], w["workstation"], w["qty_in"],
          w.get("good_qty"), w.get("scrap_qty"), w.get("downtime_seconds"), _local(w["started_at"]),
          _local(w.get("finished_at")), w.get("started_by"), w.get("edited_by_pic", False)]
         for w in db.wip.find({"tenant_openid": tenant_openid, "started_at": {"$gte": start, "$lte": end}},
                              {"_id": 0}).sort("started_at", 1)))

    counts["Scrap"] = _sheet(wb, "Scrap",
        ["Event ID", "Time", "WIP ID", "Batch ID", "WO ID", "Workstation", "Step", "Qty", "Reason", "Note",
         "Staff ID", "Source", "Edited by PIC"],
        ([e["event_id"], _local(e["ts"]), e["wip_id"], e["batch_id"], e["wo_id"], e["workstation"], e["step_no"],
          e["qty"], e["reason_code"], e.get("note", ""), e.get("staff_id"), e.get("source"), e.get("edited_by_pic", False)]
         for e in _events(db, tenant_openid, start, end, "scrap")))

    counts["Downtime"] = _sheet(wb, "Downtime",
        ["WIP ID", "Workstation", "Reason", "Stopped", "Resumed", "Minutes", "Staff ID"],
        _downtime_rows(db, tenant_openid, start, end))

    counts["Measurements"] = _sheet(wb, "Measurements",
        ["Event ID", "Time", "WIP ID", "Workstation", "Measure", "Value", "Unit", "Target", "In spec", "Staff ID"],
        ([e["event_id"], _local(e["ts"]), e["wip_id"], e["workstation"], e["measure_name"], e["value"], e["unit"],
          e.get("target"), e.get("in_spec"), e.get("staff_id")]
         for e in _events(db, tenant_openid, start, end, "measure")))

    counts["Stock movements"] = _sheet(wb, "Stock movements",
        ["Time", "Mode", "Bin", "Goods code", "Goods", "Qty", "By"],
        ([_local(r.create_time), r.mode_code, r.bin_name, r.goods_code, r.goods_desc, r.goods_qty, r.creater]
         for r in QTYRecorder.objects.filter(openid=tenant_openid, create_time__gte=_local(start),
                                             create_time__lte=_local(end)).order_by("create_time").iterator()))

    counts["Audit trail"] = _sheet(wb, "Audit trail",
        ["Time", "Who", "Role", "Action", "Entity", "Entity ID", "Before", "After", "Reason", "IP"],
        ([_local(a.ts), a.actor_name, a.actor_role, a.action, a.entity_type, a.entity_id, str(a.before),
          str(a.after), a.reason, a.ip]
         for a in AuditLog.objects.filter(tenant_openid=tenant_openid, ts__gte=_local(start), ts__lte=_local(end))
         .order_by("ts").iterator()))

    logger.info("Excel export %s..%s for %s: %s", date_from, date_to, tenant_openid, counts)
    return wb, counts


def _downtime_rows(db, tenant, start, end):
    open_stops = {}
    for e in db.step_events.find({"tenant_openid": tenant, "type": {"$in": ["stop", "resume"]},
                                  "ts": {"$gte": start, "$lte": end}}, {"_id": 0}).sort("ts", 1):
        if e["type"] == "stop":
            open_stops[e["wip_id"]] = e
        elif e["wip_id"] in open_stops:
            s = open_stops.pop(e["wip_id"])
            yield [s["wip_id"], s["workstation"], s["reason_code"], _local(s["ts"]), _local(e["ts"]),
                   round((e["ts"] - s["ts"]).total_seconds() / 60, 1), s.get("staff_id")]
    for s in open_stops.values():
        yield [s["wip_id"], s["workstation"], s["reason_code"], _local(s["ts"]), None, None, s.get("staff_id")]


def workbook_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
